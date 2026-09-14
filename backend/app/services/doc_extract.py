"""Best-effort plain-text extraction from uploaded documents.

Supports PDF (pypdf), DOCX (python-docx) and text-like files. Every path is
defensive: extraction never raises to the caller — on failure it returns an
empty string so uploads always succeed and AI features simply degrade.
"""

from __future__ import annotations

import io
import os


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        return ""
    try:
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _from_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except Exception:
        return ""
    try:
        document = docx.Document(io.BytesIO(data))
        lines = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        # include table cell text too
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines).strip()
    except Exception:
        return ""


def extract_text(filename: str, data: bytes) -> str:
    """Return extracted plain text for a supported file, else ""."""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".pdf":
        return _from_pdf(data)
    if ext in (".docx",):
        return _from_docx(data)
    if ext in (".txt", ".md", ".json", ".csv", ".rtf"):
        try:
            return data.decode("utf-8", errors="ignore").strip()
        except Exception:
            return ""
    # .doc (legacy binary) and images: not extractable here
    return ""


def looks_like_resume_text(text: str) -> bool:
    """Loose check that we actually have usable resume content."""
    if not text or len(text.strip()) < 60:
        return False
    words = text.split()
    return len(words) >= 20
